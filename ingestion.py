"""
Parses uploaded files (pdf/docx/pptx/txt) into plain text, then splits into
overlapping chunks ready for embedding.
"""
from __future__ import annotations
import io
from dataclasses import dataclass

from pypdf import PdfReader
from docx import Document as DocxDocument
from pptx import Presentation

from config import CHUNK_SIZE, CHUNK_OVERLAP


@dataclass
class Chunk:
    text: str
    source: str          # original filename
    doc_type: str        # "pricing", "case_study", "template", "other" (user-tagged)
    chunk_index: int


def extract_text(file_bytes: bytes, filename: str) -> str:
    """Route to the right parser based on file extension."""
    ext = filename.lower().rsplit(".", 1)[-1]

    if ext == "pdf":
        reader = PdfReader(io.BytesIO(file_bytes))
        return "\n".join(page.extract_text() or "" for page in reader.pages)

    if ext == "docx":
        doc = DocxDocument(io.BytesIO(file_bytes))
        return "\n".join(p.text for p in doc.paragraphs)

    if ext == "pptx":
        prs = Presentation(io.BytesIO(file_bytes))
        lines = []
        for slide in prs.slides:
            for shape in slide.shapes:
                if hasattr(shape, "text") and shape.text:
                    lines.append(shape.text)
        return "\n".join(lines)

    if ext == "txt" or ext == "md":
        return file_bytes.decode("utf-8", errors="ignore")

    raise ValueError(f"Unsupported file type: .{ext}")


def chunk_text(text: str, chunk_size: int = CHUNK_SIZE, overlap: int = CHUNK_OVERLAP) -> list[str]:
    """Simple sliding-window chunker on character count. Good enough for a
    resume project; swap for a token-aware splitter (e.g. langchain's
    RecursiveCharacterTextSplitter) if you want to be fancier."""
    text = " ".join(text.split())  # collapse whitespace/newlines
    if not text:
        return []

    chunks = []
    start = 0
    while start < len(text):
        end = start + chunk_size
        chunks.append(text[start:end])
        if end >= len(text):
            break
        start = end - overlap
    return chunks


def process_file(file_bytes: bytes, filename: str, doc_type: str) -> list[Chunk]:
    """Full pipeline: bytes -> text -> chunks -> Chunk objects with metadata."""
    text = extract_text(file_bytes, filename)
    raw_chunks = chunk_text(text)
    return [
        Chunk(text=c, source=filename, doc_type=doc_type, chunk_index=i)
        for i, c in enumerate(raw_chunks)
    ]
